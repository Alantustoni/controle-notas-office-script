"""Gera as planilhas do repositório (dados 100% fictícios).

    python tools/gerar_planilhas.py

- modelo/Controle_Notas_modelo.xlsx  → vazia, pronta para usar com o script
- modelo/Controle_Notas_exemplo.xlsx → preenchida com a saída real do script
  rodando sobre e-mails fictícios (modelo/exemplo.json, gerado por
  tools/gerar_exemplo.mjs)

Requer: pip install openpyxl
"""
import json
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

RAIZ = Path(__file__).resolve().parent.parent
PASTA = RAIZ / "modelo"

# Cabeçalhos exigidos pelo script (não podem ser renomeados)
CAB_NOTAS = ["NF", "Empresa", "Fornecedor", "FOR", "AF / Pedido", "Vencimento", "Comunicação",
             "Recebido em", "Conferencia", "Possivel_Vinculo", "Observacoes", "Valor_NF_Informado",
             "Valor_AF_Informado", "Diferenca_Informada", "Liquido_Informado", "ID_Email",
             "ID_Registro", "Chave_Provisoria"]
CAB_EMAILS = ["ID_Email", "Message_ID", "Internet_Message_ID", "Recebido_UTC_3",
              "Recebido_UTC_Original", "Remetente", "Assunto", "Referencias", "Pendencias",
              "Status_Captura", "Tem_Anexos", "Link_Original", "Trecho_Corpo"]
CAB_REVISAO = ["ID_Revisao", "ID_Email", "NF", "Tipo", "Detalhe", "Situacao"]

EMPRESAS = ["ALFAMETAL", "BETAVIARIA"]      # igual a EMPRESAS no .ts
LINHA_CAB_NOTAS = 12                        # tbNotas começa abaixo do painel

AZUL = "1F4E79"
TITULO = Font(bold=True, size=16, color=AZUL)
ROTULO = Font(bold=True, size=10, color="595959")
NUMERO = Font(bold=True, size=22, color=AZUL)
NOTA = Font(italic=True, size=9, color="7F7F7F")
DUPLICADO = PatternFill("solid", fgColor="FFF2CC")


def formula_contador(empresa):
    return ('=IFERROR(SUMPRODUCT(--(tbNotas[Empresa]="' + empresa + '"),--(tbNotas[NF]<>""),'
            'LEN(tbNotas[NF])-LEN(SUBSTITUTE(tbNotas[NF],"/",""))+1),0)')


def tabela(ws, nome, cab, dados, linha0=1, larguras=None):
    for j, c in enumerate(cab, 1):
        ws.cell(row=linha0, column=j, value=c)
    if not dados:
        dados = [[None] * len(cab)]               # o script remove esta linha vazia
    for i, linha in enumerate(dados, linha0 + 1):
        for j, v in enumerate(linha, 1):
            ws.cell(row=i, column=j, value=None if v == "" else v)
    fim = linha0 + len(dados)
    ref = f"A{linha0}:{get_column_letter(len(cab))}{fim}"
    t = Table(displayName=nome, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)
    for j, w in enumerate(larguras or [], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    return linha0 + 1, fim


def formatar_colunas(ws, l1, l2, formatos):
    for col, fmt in formatos.items():
        for r in range(l1, l2 + 1):
            ws.cell(row=r, column=col).number_format = fmt


def montar(dados, destino):
    wb = Workbook()

    # ── Notas recebidas (painel + tbNotas) ──
    ws = wb.active
    ws.title = "Notas recebidas"
    ws["B2"] = "Controle de Notas — Compras"
    ws["B2"].font = TITULO
    ws["B3"] = "Leitor v1.3 • Campos ausentes ficam em branco. Pendências na aba Revisao."
    ws["B3"].font = NOTA
    for (rot, num), empresa in zip([("B5", "B6"), ("F5", "F6")], EMPRESAS):
        ws[rot] = f"Referências {empresa}"
        ws[rot].font = ROTULO
        ws[num] = formula_contador(empresa)
        ws[num].font = NUMERO
        ws[num].alignment = Alignment(horizontal="left")
    ws["B10"] = ("Totais de referências (não NFs únicas). Barras somam cada NF; "
                 "sem empresa/NF não entra nos totais.")
    ws["B10"].font = NOTA
    larg = [12, 13, 30, 8, 12, 12, 18, 16, 20, 30, 45, 12, 12, 12, 12, 18, 22, 30]
    l1, l2 = tabela(ws, "tbNotas", CAB_NOTAS, dados["tbNotas"], LINHA_CAB_NOTAS, larg)
    formatar_colunas(ws, l1, l2, {6: "dd/mm/yyyy", 8: "dd/mm/yyyy hh:mm",
                                  **{c: "#,##0.00" for c in (12, 13, 14, 15)}})
    # mesmo destaque que o script aplica no Excel: NF + fornecedor repetidos
    grupos = defaultdict(list)
    for i, r in enumerate(dados["tbNotas"]):
        if r[0] and r[2]:
            grupos[(str(r[0]).upper(), str(r[2]).upper())].append(i)
    for idx in grupos.values():
        if len(idx) > 1:
            for i in idx:
                for j in range(1, len(CAB_NOTAS) + 1):
                    c = ws.cell(row=l1 + i, column=j)
                    c.fill = DUPLICADO
                    c.font = Font(bold=True, color="7F6000")
    ws.freeze_panes = f"A{LINHA_CAB_NOTAS + 1}"

    # ── Emails ──
    we = wb.create_sheet("Emails")
    l1, l2 = tabela(we, "tbEmails", CAB_EMAILS, dados["tbEmails"],
                    larguras=[22, 18, 20, 16, 22, 26, 30, 11, 11, 15, 22, 20, 50])
    formatar_colunas(we, l1, l2, {4: "dd/mm/yyyy hh:mm"})

    # ── Revisao ──
    wr = wb.create_sheet("Revisao")
    tabela(wr, "tbRevisao", CAB_REVISAO, dados["tbRevisao"], larguras=[24, 22, 12, 32, 80, 12])

    # ── Guia ──
    wg = wb.create_sheet("Guia")
    linhas = [
        ("Como esta planilha funciona", TITULO),
        ("", None),
        ("1. O Power Automate chama o Office Script a cada e-mail de Compras.", None),
        ("2. O script lê a tabela do e-mail e grava cada NF em tbNotas (novas no topo).", None),
        ("3. Cada e-mail fica registrado em tbEmails; o mesmo e-mail nunca entra duas vezes.", None),
        ("4. O que o script não conseguiu ler com segurança vai para tbRevisao.", None),
        ("5. Linhas amarelas = mesma NF e fornecedor em mais de um registro: conferir.", None),
        ("", None),
        ("Não renomeie as tabelas nem as colunas: o script confere os cabeçalhos.", ROTULO),
        ("Os dados desta planilha são fictícios.", NOTA),
    ]
    for i, (texto, fonte) in enumerate(linhas, 2):
        c = wg.cell(row=i, column=2, value=texto)
        if fonte:
            c.font = fonte
    wg.column_dimensions["B"].width = 90

    wb.save(destino)
    print("gravado:", destino.relative_to(RAIZ))


if __name__ == "__main__":
    PASTA.mkdir(exist_ok=True)
    vazio = {"tbNotas": [], "tbEmails": [], "tbRevisao": []}
    montar(vazio, PASTA / "Controle_Notas_modelo.xlsx")
    exemplo = json.loads((PASTA / "exemplo.json").read_text(encoding="utf-8"))
    montar({k: exemplo[k]["dados"] for k in vazio}, PASTA / "Controle_Notas_exemplo.xlsx")
